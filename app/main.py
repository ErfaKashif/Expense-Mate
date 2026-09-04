import docx
from bs4 import BeautifulSoup
import markdown

# Read the SRS Markdown file
with open("../docs/Phase1_SRS.md", "r", encoding="utf-8") as f:
    md_content = f.read()

# Convert Markdown to HTML
html_content = markdown.markdown(md_content, extensions=['tables'])

# Create Word Document
doc = docx.Document()
soup = BeautifulSoup(html_content, 'html.parser')

for element in soup.children:
    if element.name == 'h1':
        doc.add_heading(element.text, level=1)
    elif element.name == 'h2':
        doc.add_heading(element.text, level=2)
    elif element.name == 'h3':
        doc.add_heading(element.text, level=3)
    elif element.name == 'p':
        doc.add_paragraph(element.text)
    elif element.name == 'ul':
        for li in element.find_all('li'):
            doc.add_paragraph(li.text, style='List Bullet')
    elif element.name == 'table':
        rows = element.find_all('tr')
        table = doc.add_table(rows=len(rows), cols=len(rows[0].find_all(['th', 'td'])))
        table.style = 'Table Grid'
        for i, row in enumerate(rows):
            cols = row.find_all(['th', 'td'])
            for j, col in enumerate(cols):
                table.cell(i, j).text = col.text.strip()

# Save Word file
doc.save("SRS_ExpenseMate.docx")
print("Successfully generated SRS_ExpenseMate.docx")